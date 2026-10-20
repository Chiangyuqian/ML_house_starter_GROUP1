from pathlib import Path
import json
import time
import joblib
import pandas as pd

from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from sklearn.compose import ColumnTransformer
from sklearn.ensemble import RandomForestRegressor
from sklearn.impute import SimpleImputer
from sklearn.linear_model import LinearRegression
from sklearn.metrics import mean_absolute_error, mean_squared_error, r2_score
from sklearn.model_selection import train_test_split
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import OneHotEncoder

ROOT = Path(__file__).parent
DATA = ROOT / 'houses_demo.csv'
OUT = ROOT / 'outputs'
OUT.mkdir(exist_ok=True)

TARGET = '總價萬元'
NUM = ['坪數', '屋齡', '距捷運公里', '房間數']
CAT = ['區域']

# 讀取資料
df = pd.read_csv(DATA, encoding='utf-8-sig')
df = df.dropna(subset=[TARGET])

X = df[NUM + CAT]
y = df[TARGET]

# 分割訓練與測試資料
X_train, X_test, y_train, y_test = train_test_split(
    X, y,
    test_size=0.2,
    random_state=42
)

# 建立 Pipeline
def make_pipeline(regressor):
    preprocess = ColumnTransformer([
        ('num', SimpleImputer(strategy='median'), NUM),
        ('cat', Pipeline([
            ('fill', SimpleImputer(strategy='most_frequent')),
            ('onehot', OneHotEncoder(handle_unknown='ignore'))
        ]), CAT)
    ])

    return Pipeline([
        ('preprocess', preprocess),
        ('model', regressor)
    ])


# 定義模型
models = [
    ('線性迴歸', LinearRegression()),
    ('隨機森林', RandomForestRegressor(
        n_estimators=150,
        random_state=42,
        min_samples_leaf=2
    ))
]

scores = {}
records = []
best_model = None
best_model_name = None
best_mae = float('inf')
best_index = None

print('========== 開始模型訓練 ==========')

for name, regressor in models:

    pipeline = make_pipeline(regressor)

    # 記錄訓練時間
    start_time = time.perf_counter()

    pipeline.fit(X_train, y_train)

    training_time = time.perf_counter() - start_time

    # 預測測試資料
    pred = pipeline.predict(X_test)

    # 評估指標
    mae = float(mean_absolute_error(y_test, pred))
    rmse = float(mean_squared_error(y_test, pred) ** 0.5)
    r2 = float(r2_score(y_test, pred))

    scores[name] = {
        'MAE_萬元': round(mae, 2),
        'RMSE_萬元': round(rmse, 2),
        'R2': round(r2, 4)
    }

    records.append({
        '模型名稱': name,
        '訓練筆數': len(X_train),
        '測試筆數': len(X_test),
        '訓練時間(秒)': round(training_time, 4),
        'MAE(萬元)': round(mae, 2),
        'RMSE(萬元)': round(rmse, 2),
        'R²': round(r2, 4),
        '最佳MAE(萬元)': None,
        '最終選用模型': None
    })

    print(f'\n模型：{name}')
    print(f'訓練時間：{training_time:.4f} 秒')
    print(f'MAE：{mae:.2f} 萬元')
    print(f'RMSE：{rmse:.2f} 萬元')
    print(f'R²：{r2:.4f}')

    # 選擇 MAE 最小的模型
    if mae < best_mae:
        best_mae = mae
        best_model = pipeline
        best_model_name = name
        best_index = len(records) - 1


# 在最佳模型那列填入結果
records[best_index]['最佳MAE(萬元)'] = round(best_mae, 2)
records[best_index]['最終選用模型'] = best_model_name

# 儲存最佳模型
joblib.dump(
    best_model,
    OUT / 'house_price_model.joblib'
)

# 儲存 JSON
(OUT / 'metrics.json').write_text(
    json.dumps(scores, ensure_ascii=False, indent=2),
    encoding='utf-8'
)

# 建立 Excel 資料
result_df = pd.DataFrame(records)

excel_path = OUT / '模型訓練結果.xlsx'

with pd.ExcelWriter(excel_path, engine='openpyxl') as writer:

    result_df.to_excel(
        writer,
        sheet_name='模型訓練結果',
        index=False
    )

    ws = writer.sheets['模型訓練結果']

    # 表頭樣式
    for cell in ws[1]:
        cell.fill = PatternFill(
            'solid',
            fgColor='1F4E78'
        )
        cell.font = Font(
            bold=True,
            color='FFFFFF'
        )
        cell.alignment = Alignment(
            horizontal='center'
        )

    # 自動調整欄寬
    for col in ws.columns:
        max_length = max(
            len(str(cell.value))
            if cell.value is not None else 0
            for cell in col
        )

        letter = get_column_letter(col[0].column)
        ws.column_dimensions[letter].width = max_length * 2 + 4

    # 置中與最佳模型標示
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(
                horizontal='center'
            )

        if row[0].value == best_model_name:
            for cell in row:
                cell.fill = PatternFill(
                    'solid',
                    fgColor='E2F0D9'
                )

    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions


print('\n========== 模型訓練完成 ==========')
print(f'最終選用模型：{best_model_name}')
print(f'最佳 MAE：{best_mae:.2f} 萬元')
print(f'模型檔案：{OUT / "house_price_model.joblib"}')
print(f'Excel 報表：{excel_path}')
