from pathlib import Path
import joblib
import pandas as pd

from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).parent
OUT = ROOT / 'outputs'
OUT.mkdir(exist_ok=True)

# 載入模型
model = joblib.load(
    OUT / 'house_price_model.joblib'
)

# 取得模型名稱
model_type = type(model.named_steps['model']).__name__

model_names = {
    'LinearRegression': '線性迴歸',
    'RandomForestRegressor': '隨機森林'
}

model_name = model_names.get(model_type, model_type)

print('========== 房價預測 ==========')
print(f'使用模型：{model_name}')

# 建立預測房屋資料
new_house = pd.DataFrame([
    {
        '坪數': 35,
        '屋齡': 10,
        '距捷運公里': 1.2,
        '房間數': 3,
        '區域': '住宅區'
    }
])

# 進行預測
predictions = model.predict(new_house)

# 建立預測結果表格
result_df = new_house.copy()

result_df['預測總價(萬元)'] = [
    round(float(p), 2) for p in predictions
]

result_df['使用模型'] = model_name

# 印出結果
print('\n========== 預測結果 ==========')

for price in predictions:
    print(f'預測總價：{float(price):,.2f} 萬元')

# 匯出 Excel
excel_path = OUT / '房價預測結果.xlsx'

with pd.ExcelWriter(excel_path, engine='openpyxl') as writer:

    result_df.to_excel(
        writer,
        sheet_name='房價預測結果',
        index=False
    )

    ws = writer.sheets['房價預測結果']

    # 表頭樣式
    for cell in ws[1]:
        cell.fill = PatternFill(
            'solid',
            fgColor='1F4E78'
        )
        cell.font = Font(
            bold=True,
            color='FFFFFF'
        )
        cell.alignment = Alignment(
            horizontal='center'
        )

    # 自動調整欄寬
    for col in ws.columns:
        max_length = max(
            len(str(cell.value))
            if cell.value is not None else 0
            for cell in col
        )

        letter = get_column_letter(col[0].column)
        ws.column_dimensions[letter].width = max_length * 2 + 4

    # 設定預測價格欄位格式
    price_col = result_df.columns.get_loc('預測總價(萬元)') + 1

    for row in ws.iter_rows(min_row=2):

        for cell in row:
            cell.alignment = Alignment(
                horizontal='center'
            )

        row[price_col - 1].number_format = '#,##0.00'

        row[price_col - 1].fill = PatternFill(
            'solid',
            fgColor='E2F0D9'
        )

    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions


print(f'\nExcel 已儲存：{excel_path}')
print('注意：資料為模擬資料，僅供教學使用')
